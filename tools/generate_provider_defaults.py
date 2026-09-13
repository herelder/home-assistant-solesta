#!/usr/bin/env python3
"""Validate ProviderDefaults.xlsx and generate ``provider_defaults.py``."""

from __future__ import annotations

import argparse
from collections.abc import Mapping, Sequence
import hashlib
import json
import math
from pathlib import Path
import re
import tempfile
from typing import Any
from urllib.parse import urlsplit


PRESET_ID = "solesta"
PRESET_REVISION = 3
EXPECTED_FLOW_HEADERS = (
    "preset_id",
    "preset_revision",
    "flow_name",
    "step_index",
    "op",
    "base_url",
    "total_timeout_seconds",
    "method",
    "path_template",
    "json_payload",
    "success_statuses",
    "extract",
    "assert_equals",
)
EXPECTED_KEY_HEADERS = (
    "preset_id",
    "preset_revision",
    "available_key",
)
REQUIRED_FLOWS = (
    "setup_httpsession_flow",
    "connect_flow",
    "disconnect_flow",
    "profile_flow",
    "systems_flow",
    "system_status_flow",
)
PLACEHOLDER_WHITELIST = {
    "setup_httpsession_flow": frozenset(),
    "connect_flow": frozenset({"userid", "password"}),
    "disconnect_flow": frozenset(),
    "profile_flow": frozenset(),
    "systems_flow": frozenset(),
    "system_status_flow": frozenset({"source_system_id"}),
}
PLACEHOLDER = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")
TOKEN = re.compile(r"^\{([A-Za-z_][A-Za-z0-9_]*)\}$")
IDENTIFIER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
CONTROL_OR_SPACE = re.compile(r"[\x00-\x20\x7f]")
EXPECTED_BASE_URL = "https://dashboard.solesta.nl"
MAX_REQUEST_JSON_NODES = 256
MAX_REQUEST_JSON_DEPTH = 8
MAX_TEXT_BYTES = 64 * 1024


class ProviderWorkbookError(ValueError):
    """ProviderDefaults.xlsx violates the release contract."""


def _fail(message: str) -> None:
    raise ProviderWorkbookError(f"Provider defaults workbook: {message}")


def _is_integer(value: object) -> bool:
    return isinstance(value, int) and not isinstance(value, bool)


def _non_empty_rows(rows: Sequence[Sequence[object]]) -> list[tuple[object, ...]]:
    return [
        tuple(row)
        for row in rows
        if any(value is not None and value != "" for value in row)
    ]


def _parse_json(value: object, field: str, row_number: int) -> Any:
    if value is None or value == "":
        return None
    if not isinstance(value, str):
        _fail(f"{field} at row {row_number} must contain strict JSON text")
    try:
        return json.loads(value)
    except json.JSONDecodeError as err:
        _fail(f"{field} at row {row_number} is invalid JSON: {err.msg}")


def _tokens_in(value: object) -> set[str]:
    tokens: set[str] = set()

    def visit(item: object) -> None:
        if isinstance(item, str):
            tokens.update(PLACEHOLDER.findall(item))
        elif isinstance(item, list):
            for child in item:
                visit(child)
        elif isinstance(item, Mapping):
            for child in item.values():
                visit(child)

    visit(value)
    return tokens


def _has_value(value: object) -> bool:
    return value is not None and value != ""


def _canonical_https_origin(value: object, row_number: int) -> str:
    if not isinstance(value, str):
        _fail(f"base_url at row {row_number} must be a string")
    candidate = value.strip()
    if not candidate or CONTROL_OR_SPACE.search(candidate) or "\\" in candidate:
        _fail(f"base_url at row {row_number} is not a safe HTTPS origin")
    try:
        parsed = urlsplit(candidate)
        port = parsed.port
    except ValueError:
        _fail(f"base_url at row {row_number} is not a valid URL")
    if (
        parsed.scheme.lower() != "https"
        or parsed.username is not None
        or parsed.password is not None
        or not parsed.hostname
        or parsed.path not in ("", "/")
        or parsed.query
        or parsed.fragment
        or (port is not None and not 1 <= port <= 65535)
    ):
        _fail(f"base_url at row {row_number} is not a safe HTTPS origin")
    try:
        host = parsed.hostname.encode("idna").decode("ascii").lower()
    except UnicodeError:
        _fail(f"base_url hostname at row {row_number} is invalid")
    if ":" in host:
        host = f"[{host}]"
    port_suffix = "" if port in (None, 443) else f":{port}"
    return f"https://{host}{port_suffix}"


def _validate_json_value(
    value: Any,
    *,
    row_number: int,
    depth: int = 0,
    nodes: list[int] | None = None,
) -> None:
    if nodes is None:
        nodes = [0]
    nodes[0] += 1
    if nodes[0] > MAX_REQUEST_JSON_NODES or depth > MAX_REQUEST_JSON_DEPTH:
        _fail(f"json_payload at row {row_number} exceeds its complexity limit")
    if isinstance(value, str):
        if len(value.encode("utf-8")) > MAX_TEXT_BYTES:
            _fail(f"json_payload string at row {row_number} is too long")
        return
    if value is None or isinstance(value, (bool, int)):
        return
    if isinstance(value, float):
        if not math.isfinite(value):
            _fail(f"json_payload at row {row_number} contains a non-finite number")
        return
    if isinstance(value, list):
        for item in value:
            _validate_json_value(
                item,
                row_number=row_number,
                depth=depth + 1,
                nodes=nodes,
            )
        return
    if isinstance(value, Mapping):
        for key, item in value.items():
            if not isinstance(key, str) or len(key.encode("utf-8")) > MAX_TEXT_BYTES:
                _fail(f"json_payload at row {row_number} has an invalid key")
            _validate_json_value(
                item,
                row_number=row_number,
                depth=depth + 1,
                nodes=nodes,
            )
        return
    _fail(f"json_payload at row {row_number} contains an unsupported value")


def _validate_path_template(path: object, flow_name: str, row_number: int) -> str:
    if (
        not isinstance(path, str)
        or len(path) > 1024
        or not path.startswith("/")
        or path.startswith("//")
        or "\\" in path
        or CONTROL_OR_SPACE.search(path)
    ):
        _fail(f"unsafe path_template at row {row_number}")
    parsed = urlsplit(path)
    if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment:
        _fail(f"unsafe path_template at row {row_number}")
    allowed = PLACEHOLDER_WHITELIST[flow_name]
    for segment in path.split("/")[1:]:
        if segment in ("", ".", ".."):
            _fail(f"unsafe path segment at row {row_number}")
        match = TOKEN.fullmatch(segment)
        if match:
            if match.group(1) not in allowed:
                _fail(f"placeholder {segment} is not allowed in {flow_name}")
        elif "{" in segment or "}" in segment:
            _fail(f"placeholder must occupy a full path segment at row {row_number}")
    return path


def _validate_json_placeholders(value: Any, flow_name: str, row_number: int) -> None:
    if isinstance(value, str):
        match = TOKEN.fullmatch(value)
        if match:
            if match.group(1) not in PLACEHOLDER_WHITELIST[flow_name]:
                _fail(f"placeholder {value} is not allowed in {flow_name}")
        elif "{" in value or "}" in value:
            _fail(
                f"JSON placeholder must occupy a complete string at row {row_number}"
            )
        return
    if isinstance(value, list):
        for item in value:
            _validate_json_placeholders(item, flow_name, row_number)
        return
    if isinstance(value, Mapping):
        for item in value.values():
            _validate_json_placeholders(item, flow_name, row_number)


def _validate_response_contract(
    success_statuses: Any,
    extract: Any,
    assert_equals: Any,
    row_number: int,
) -> None:
    if (
        not isinstance(success_statuses, list)
        or not 1 <= len(success_statuses) <= 16
        or any(
            not _is_integer(status) or not 200 <= status <= 299
            for status in success_statuses
        )
        or len(success_statuses) != len(set(success_statuses))
    ):
        _fail(f"invalid success_statuses at row {row_number}")
    if (
        not isinstance(extract, list)
        or len(extract) > 64
        or any(not isinstance(selector, str) or not selector for selector in extract)
        or len(extract) != len(set(extract))
        or ("*" in extract and extract != ["*"])
    ):
        _fail(f"invalid extract at row {row_number}")
    if assert_equals is None:
        return
    if not isinstance(assert_equals, Mapping) or not 1 <= len(assert_equals) <= 16:
        _fail(f"assert_equals at row {row_number} must be a non-empty object")
    for key, expected in assert_equals.items():
        if (
            not isinstance(key, str)
            or not key
            or len(key.encode("utf-8")) > MAX_TEXT_BYTES
        ):
            _fail(f"assert_equals at row {row_number} contains an invalid key")
        if isinstance(expected, str):
            if len(expected.encode("utf-8")) > MAX_TEXT_BYTES:
                _fail(f"assert_equals string at row {row_number} is too long")
        elif isinstance(expected, float):
            if not math.isfinite(expected):
                _fail(f"assert_equals at row {row_number} contains a non-finite number")
        elif expected is not None and not isinstance(expected, (bool, int)):
            _fail(f"assert_equals at row {row_number} must contain JSON scalars")


def _load_source_rows(path: Path) -> tuple[list[tuple[object, ...]], list[tuple[object, ...]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as err:  # pragma: no cover - environment-specific
        raise RuntimeError("generate_provider_defaults.py requires openpyxl") from err

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        required_sheets = {"ProviderRESTFlows", "ProviderAvailableKeys"}
        missing_sheets = required_sheets - set(workbook.sheetnames)
        if missing_sheets:
            _fail("missing worksheets: " + ", ".join(sorted(missing_sheets)))

        flow_values = list(
            workbook["ProviderRESTFlows"].iter_rows(
                min_row=4, max_row=200, min_col=1, max_col=13, values_only=True
            )
        )
        key_values = list(
            workbook["ProviderAvailableKeys"].iter_rows(
                min_row=4, max_row=200, min_col=1, max_col=3, values_only=True
            )
        )
    finally:
        workbook.close()

    if not flow_values or tuple(flow_values[0]) != EXPECTED_FLOW_HEADERS:
        _fail("ProviderRESTFlows headers do not match the release schema")
    if not key_values or tuple(key_values[0]) != EXPECTED_KEY_HEADERS:
        _fail("ProviderAvailableKeys headers do not match the release schema")

    flow_rows = _non_empty_rows(flow_values[1:])
    key_rows = _non_empty_rows(key_values[1:])
    if not flow_rows:
        _fail("ProviderRESTFlows has no data rows")
    if not key_rows:
        _fail(
            "ProviderAvailableKeys must identify the preset even when its key list is empty"
        )
    return flow_rows, key_rows


def load_provider_default(path: Path) -> dict[str, Any]:
    """Return the validated provider structure from the canonical workbook."""
    flow_rows, key_rows = _load_source_rows(path)
    preset_id, preset_revision = flow_rows[0][0:2]
    if preset_id != PRESET_ID or preset_revision != PRESET_REVISION:
        _fail(
            f"expected preset_id {PRESET_ID!r} and preset_revision {PRESET_REVISION}"
        )

    flows: dict[str, list[dict[str, Any]]] = {
        name: [] for name in REQUIRED_FLOWS
    }
    for index, row in enumerate(flow_rows):
        row_number = index + 5
        (
            row_preset,
            row_revision,
            flow_name,
            step_index,
            operation,
            base_url,
            timeout,
            method,
            path_template,
            json_text,
            statuses_text,
            extract_text,
            assert_text,
        ) = row
        if row_preset != preset_id or row_revision != preset_revision:
            _fail(f"preset mismatch at row {row_number}")
        if flow_name not in flows:
            _fail(f"unsupported flow_name at row {row_number}")
        if not _is_integer(step_index) or step_index != len(flows[flow_name]):
            _fail(f"step_index is not contiguous at row {row_number}")

        if operation == "create_httpsession":
            if (
                flow_name != "setup_httpsession_flow"
                or _canonical_https_origin(base_url, row_number)
                != EXPECTED_BASE_URL
            ):
                _fail(f"unsafe session setup at row {row_number}")
            if not _is_integer(timeout) or not 1 <= timeout <= 60:
                _fail(f"invalid timeout at row {row_number}")
            if any(
                _has_value(value)
                for value in (
                    method,
                    path_template,
                    json_text,
                    statuses_text,
                    extract_text,
                    assert_text,
                )
            ):
                _fail(f"session setup has request fields at row {row_number}")
            step: dict[str, Any] = {
                "op": operation,
                "base_url": EXPECTED_BASE_URL,
                "total_timeout_seconds": timeout,
            }
        elif operation == "request":
            if _has_value(base_url) or _has_value(timeout):
                _fail(f"request has session fields at row {row_number}")
            if method not in {"GET", "POST"}:
                _fail(f"unsupported method at row {row_number}")
            if (
                flow_name
                in {"profile_flow", "systems_flow", "system_status_flow"}
                and method != "GET"
            ):
                _fail(f"{flow_name} may contain only GET steps")
            path_template = _validate_path_template(
                path_template,
                flow_name,
                row_number,
            )
            payload = _parse_json(json_text, "json_payload", row_number)
            success_statuses = _parse_json(
                statuses_text, "success_statuses", row_number
            )
            extract = _parse_json(extract_text, "extract", row_number)
            assert_equals = _parse_json(assert_text, "assert_equals", row_number)
            if method == "GET" and payload is not None:
                _fail(f"GET cannot have json_payload at row {row_number}")
            if payload is not None:
                if not isinstance(payload, Mapping):
                    _fail(f"POST json_payload at row {row_number} must be an object")
                _validate_json_value(payload, row_number=row_number)
                _validate_json_placeholders(payload, flow_name, row_number)
            _validate_response_contract(
                success_statuses,
                extract,
                assert_equals,
                row_number,
            )
            step = {
                "op": operation,
                "method": method,
                "path_template": path_template,
            }
            if payload is not None:
                step["json"] = payload
            step["response"] = {
                "success_statuses": success_statuses,
                "extract": extract,
            }
            if assert_equals is not None:
                step["response"]["assert_equals"] = assert_equals
        else:
            _fail(f"unsupported op at row {row_number}")

        allowed = PLACEHOLDER_WHITELIST[flow_name]
        for token in _tokens_in(step):
            if token not in allowed:
                _fail(f"placeholder {{{token}}} is not allowed in {flow_name}")
        flows[flow_name].append(step)

    for flow_name in REQUIRED_FLOWS:
        if len(flows[flow_name]) != 1:
            _fail(f"the Solesta preset requires exactly one {flow_name} step")
    if flows["connect_flow"][0]["response"].get("assert_equals") != {"success": True}:
        _fail("the Solesta login step must assert success=true")

    available_keys: list[str] = []
    for index, row in enumerate(key_rows):
        row_preset, row_revision, available_key = row
        if row_preset != preset_id or row_revision != preset_revision:
            _fail(f"available-key preset mismatch at row {index + 5}")
        if available_key is None or available_key == "":
            continue
        if not isinstance(available_key, str) or not IDENTIFIER.fullmatch(
            available_key
        ):
            _fail(f"invalid available_key at row {index + 5}")
        available_keys.append(available_key)
    if len(set(available_keys)) != len(available_keys):
        _fail("available_key values must be unique")
    available_keys.sort()

    setup = flows.pop("setup_httpsession_flow")[0]
    runtime_flows = {
        flow_name: {
            "bindings": sorted(PLACEHOLDER_WHITELIST[flow_name]),
            "steps": flows[flow_name],
        }
        for flow_name in REQUIRED_FLOWS[1:]
    }
    return {
        "schema_version": 1,
        "preset_id": preset_id,
        "preset_revision": preset_revision,
        "protocol": "web_api",
        "base_url": setup["base_url"],
        "total_timeout_seconds": setup["total_timeout_seconds"],
        "available_keys": available_keys,
        "flows": runtime_flows,
    }


def _python_literal(value: object, indent: int = 0) -> str:
    pad = " " * indent
    next_pad = " " * (indent + 4)
    if value is None:
        return "None"
    if value is True:
        return "True"
    if value is False:
        return "False"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, str):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, list):
        if not value:
            return "[]"
        children = "\n".join(
            f"{next_pad}{_python_literal(item, indent + 4)}," for item in value
        )
        return f"[\n{children}\n{pad}]"
    if isinstance(value, Mapping):
        if not value:
            return "{}"
        children = "\n".join(
            f"{next_pad}{json.dumps(key)}: {_python_literal(item, indent + 4)},"
            for key, item in value.items()
        )
        return f"{{\n{children}\n{pad}}}"
    raise TypeError(f"unsupported generated value: {type(value).__name__}")


def generate_text(
    workbook_path: Path,
    *,
    generator_path: Path | None = None,
) -> str:
    """Validate the canonical input and return deterministic Python source."""
    provider_default = load_provider_default(workbook_path)
    actual_generator_path = generator_path or Path(__file__)
    source_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    generator_hash = hashlib.sha256(actual_generator_path.read_bytes()).hexdigest()
    return f'''"""Generated provider defaults; do not edit manually.

Source: tools/sources/{workbook_path.name}
Source SHA-256: {source_hash}
Generator: tools/generate_provider_defaults.py
Generator SHA-256: {generator_hash}
Preset revision: {PRESET_REVISION}
"""

from __future__ import annotations

from collections.abc import Mapping
from types import MappingProxyType
from typing import Any, Final


def _freeze(value: Any) -> Any:
    """Return an immutable tree built only from validated generated data."""
    if isinstance(value, dict):
        return MappingProxyType({{key: _freeze(item) for key, item in value.items()}})
    if isinstance(value, list):
        return tuple(_freeze(item) for item in value)
    return value


SOLESTA_WEB_API_DEFAULT: Final[Mapping[str, Any]] = _freeze(
    {_python_literal(provider_default, 4)}
)


PROVIDER_DEFAULTS: Final[Mapping[str, Mapping[str, Any]]] = MappingProxyType({{
    "solesta": SOLESTA_WEB_API_DEFAULT,
}})
'''


def generate(
    workbook_path: Path,
    output_path: Path,
    *,
    generator_path: Path | None = None,
) -> None:
    """Validate first and atomically replace the generated output."""
    rendered = generate_text(workbook_path, generator_path=generator_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8",
            dir=output_path.parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            handle.write(rendered)
            temporary_path = Path(handle.name)
        temporary_path.replace(output_path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def main(argv: Sequence[str] | None = None) -> int:
    tools_dir = Path(__file__).resolve().parent
    integration_dir = tools_dir.parent / "custom_components" / "solesta"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=tools_dir / "sources" / "ProviderDefaults.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=integration_dir / "provider_defaults.py",
    )
    args = parser.parse_args(argv)
    try:
        generate(args.workbook, args.output)
    except (OSError, RuntimeError, ProviderWorkbookError) as err:
        parser.error(str(err))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
