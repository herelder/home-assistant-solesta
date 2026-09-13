#!/usr/bin/env python3
"""Validate an EntityDescriptions workbook and generate ``thetruth.py``.

The generator intentionally validates the complete workbook before opening the
output file. It uses the integration's production expression compiler rather
than maintaining a second, weaker expression parser.
"""

from __future__ import annotations

import argparse
import ast
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import hashlib
import importlib.util
from io import StringIO
import json
import math
from pathlib import Path
import re
import sys
import tempfile
import tokenize
from types import ModuleType
from typing import Any


SHEET_NAME = "EntityDescriptions"
IDENTIFIER = re.compile(r"^[a-z][a-z0-9_]*$")
ICON = re.compile(r"^mdi:[a-z0-9-]+$")
MAX_ROUND_DIGITS = 15
CATALOG_REVISION = 7
EXPECTED_KEYS = (
    "last_report",
    "temp_collector",
    "temp_tank_top",
    "temp_tank_bottom",
    "temp_reserved",
    "perc_pump_pwm",
    "temp_tank_average",
    "report_delay",
    "energy_tank_mj",
    "energy_tank_kwh",
    "energy_tank_showers",
    "state_pump",
    "alerts_system",
    "status_reporting",
    "volume_tank",
    "temp_coldwater",
    "energy_shower_kwh",
    "update_interval",
    "weekly_archive",
    "archive_last_success",
    "archive_status",
)

DESCRIPTION_COLUMNS = (
    "key",
    "translation_key",
    "name",
    "device_class",
    "native_unit_of_measurement",
    "state_class",
    "entity_category",
    "icon",
    "entity_registry_enabled_default",
    "entity_registry_visible_default",
    "suggested_display_precision",
    "has_entity_name",
    "native_min_value",
    "native_max_value",
    "native_step",
    "mode",
    "entity_type",
    "options",
    "default_source_expression",
)
REQUIRED_COLUMNS = frozenset(DESCRIPTION_COLUMNS)
AUDIT_COLUMNS = frozenset({"Description", "unique_id", "entity_id"})

ENTITY_TYPES = frozenset({"sensor", "binary_sensor", "number", "switch"})
STATE_CLASSES = frozenset({"measurement", "total", "total_increasing"})
ENTITY_CATEGORIES = frozenset({"config", "diagnostic"})
NUMBER_MODES = frozenset({"auto", "box", "slider"})

# These are deliberately explicit. A typo must fail generation instead of
# becoming an AttributeError only after Home Assistant imports the result.
SENSOR_DEVICE_CLASSES = frozenset(
    {
        "apparent_power", "aqi", "atmospheric_pressure", "battery", "co",
        "co2", "current", "data_rate", "data_size", "date", "distance",
        "duration", "energy", "energy_distance", "energy_storage", "enum",
        "frequency", "gas", "humidity", "illuminance", "irradiance",
        "moisture", "monetary", "nitrogen_dioxide", "nitrogen_monoxide",
        "nitrous_oxide", "ozone", "ph", "pm1", "pm10", "pm25", "power",
        "power_factor", "precipitation", "precipitation_intensity",
        "pressure", "reactive_energy", "reactive_power", "signal_strength",
        "sound_pressure", "speed", "sulphur_dioxide", "temperature",
        "timestamp", "volatile_organic_compounds",
        "volatile_organic_compounds_parts", "voltage", "volume",
        "volume_flow_rate", "volume_storage", "water", "weight",
        "wind_direction", "wind_speed",
    }
)
BINARY_SENSOR_DEVICE_CLASSES = frozenset(
    {
        "battery", "battery_charging", "carbon_monoxide", "cold",
        "connectivity", "door", "garage_door", "gas", "heat", "light",
        "lock", "moisture", "motion", "moving", "occupancy", "opening",
        "plug", "power", "presence", "problem", "running", "safety",
        "smoke", "sound", "tamper", "update", "vibration", "window",
    }
)
NUMBER_DEVICE_CLASSES = frozenset(
    {
        "apparent_power", "battery", "co", "co2", "current", "data_rate",
        "data_size", "distance", "duration", "energy", "frequency", "gas",
        "humidity", "illuminance", "irradiance", "moisture", "monetary",
        "nitrogen_dioxide", "nitrogen_monoxide", "nitrous_oxide", "ozone",
        "ph", "pm1", "pm10", "pm25", "power", "power_factor",
        "precipitation", "precipitation_intensity", "pressure",
        "reactive_energy", "reactive_power", "signal_strength",
        "sound_pressure", "speed", "sulphur_dioxide", "temperature",
        "volatile_organic_compounds", "volatile_organic_compounds_parts",
        "voltage", "volume", "volume_flow_rate", "volume_storage", "water",
        "weight", "wind_direction", "wind_speed",
    }
)


class WorkbookValidationError(ValueError):
    """EntityDescriptions contains one or more invalid rows."""


@dataclass(frozen=True, slots=True)
class EntityRow:
    """Normalized metadata for one generated Description."""

    row_number: int
    values: Mapping[str, Any]
    unique_id: str | None = None
    entity_id: str | None = None

    def __getitem__(self, key: str) -> Any:
        return self.values[key]


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _required_text(value: Any, *, row: int, column: str) -> str:
    text = _optional_text(value)
    if text is None:
        raise WorkbookValidationError(f"row {row}: {column} is required")
    return text


def _boolean(value: Any, *, row: int, column: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"true", "yes"}:
            return True
        if normalized in {"false", "no"}:
            return False
    raise WorkbookValidationError(
        f"row {row}: {column} must be true/false (or yes/no)"
    )


def _number(value: Any, *, row: int, column: str) -> int | float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise WorkbookValidationError(f"row {row}: {column} must be numeric")
    try:
        numeric = float(value)
    except (TypeError, ValueError) as err:
        raise WorkbookValidationError(
            f"row {row}: {column} must be numeric"
        ) from err
    if not math.isfinite(numeric):
        raise WorkbookValidationError(f"row {row}: {column} must be finite")
    return int(numeric) if numeric.is_integer() else numeric


def _precision(value: Any, *, row: int) -> int | None:
    numeric = _number(value, row=row, column="suggested_display_precision")
    if numeric is None:
        return None
    if not isinstance(numeric, int) or numeric < 0:
        raise WorkbookValidationError(
            f"row {row}: suggested_display_precision must be a non-negative integer"
        )
    return numeric


def _format_expression(expression: str) -> str:
    """Use one space after argument commas without changing string literals."""
    try:
        tokens = tuple(tokenize.generate_tokens(StringIO(expression).readline))
    except (IndentationError, tokenize.TokenError):
        # Syntax validation later in the generator owns malformed expressions.
        return expression

    lines = expression.splitlines(keepends=True)
    offsets: list[int] = []
    offset = 0
    for line in lines:
        offsets.append(offset)
        offset += len(line)

    edits: list[tuple[int, int]] = []
    for token in tokens:
        if token.type != tokenize.OP or token.string != ",":
            continue
        row, column = token.end
        start = offsets[row - 1] + column
        end = start
        while end < len(expression) and expression[end] in " \t":
            end += 1
        if end < len(expression) and expression[end] not in "\r\n":
            edits.append((start, end))

    for start, end in reversed(edits):
        expression = expression[:start] + " " + expression[end:]
    return expression


def _expression(value: Any, *, row: int) -> str:
    if isinstance(value, bool) or value is None:
        raise WorkbookValidationError(
            f"row {row}: default_source_expression is required"
        )
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            raise WorkbookValidationError(
                f"row {row}: default_source_expression must be finite"
            )
        return repr(value)
    return _format_expression(
        _required_text(
            value, row=row, column="default_source_expression"
        )
    )


def _options(value: Any, *, row: int) -> tuple[str, ...] | None:
    text = _optional_text(value)
    if text is None:
        return None
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as err:
        raise WorkbookValidationError(
            f"row {row}: options must be a JSON array of strings"
        ) from err
    if (
        not isinstance(parsed, list)
        or not parsed
        or any(not isinstance(item, str) or not item for item in parsed)
        or len(parsed) != len(set(parsed))
    ):
        raise WorkbookValidationError(
            f"row {row}: options must contain unique non-empty strings"
        )
    return tuple(parsed)


def normalize_rows(
    raw_rows: Iterable[tuple[int, Mapping[str, Any]]],
) -> list[EntityRow]:
    """Normalize worksheet dictionaries without hiding blank metadata."""
    result: list[EntityRow] = []
    for row_number, raw in raw_rows:
        if not any(value is not None and str(value).strip() for value in raw.values()):
            continue
        entity_type = _required_text(
            raw.get("entity_type"), row=row_number, column="entity_type"
        ).lower()
        values = {
            "key": _required_text(raw.get("key"), row=row_number, column="key"),
            "translation_key": _required_text(
                raw.get("translation_key"),
                row=row_number,
                column="translation_key",
            ),
            "name": _optional_text(raw.get("name")),
            "device_class": (
                _optional_text(raw.get("device_class")) or ""
            ).lower() or None,
            "native_unit_of_measurement": _optional_text(
                raw.get("native_unit_of_measurement")
            ),
            "state_class": (
                _optional_text(raw.get("state_class")) or ""
            ).lower() or None,
            "entity_category": (
                _optional_text(raw.get("entity_category")) or ""
            ).lower() or None,
            "icon": _optional_text(raw.get("icon")),
            "entity_registry_enabled_default": _boolean(
                raw.get("entity_registry_enabled_default"),
                row=row_number,
                column="entity_registry_enabled_default",
            ),
            "entity_registry_visible_default": _boolean(
                raw.get("entity_registry_visible_default"),
                row=row_number,
                column="entity_registry_visible_default",
            ),
            "suggested_display_precision": _precision(
                raw.get("suggested_display_precision"), row=row_number
            ),
            "has_entity_name": _boolean(
                raw.get("has_entity_name"),
                row=row_number,
                column="has_entity_name",
            ),
            "native_min_value": _number(
                raw.get("native_min_value"),
                row=row_number,
                column="native_min_value",
            ),
            "native_max_value": _number(
                raw.get("native_max_value"),
                row=row_number,
                column="native_max_value",
            ),
            "native_step": _number(
                raw.get("native_step"), row=row_number, column="native_step"
            ),
            "mode": (_optional_text(raw.get("mode")) or "").lower() or None,
            "entity_type": entity_type,
            "options": _options(raw.get("options"), row=row_number),
            "default_source_expression": _expression(
                raw.get("default_source_expression"), row=row_number
            ),
        }
        result.append(
            EntityRow(
                row_number=row_number,
                values=values,
                unique_id=_optional_text(raw.get("unique_id")),
                entity_id=_optional_text(raw.get("entity_id")),
            )
        )
    if not result:
        raise WorkbookValidationError("EntityDescriptions contains no entities")
    return result


def _ensure_unique(rows: Sequence[EntityRow], attribute: str) -> None:
    seen: dict[str, int] = {}
    for row in rows:
        value = row[attribute] if attribute in row.values else getattr(row, attribute)
        if value is None:
            continue
        if value in seen:
            raise WorkbookValidationError(
                f"row {row.row_number}: duplicate {attribute} {value!r}; "
                f"first used on row {seen[value]}"
            )
        seen[value] = row.row_number


def _validate_number_metadata(row: EntityRow) -> None:
    minimum = row["native_min_value"]
    maximum = row["native_max_value"]
    step = row["native_step"]
    mode = row["mode"]
    if any(value is None for value in (minimum, maximum, step, mode)):
        raise WorkbookValidationError(
            f"row {row.row_number}: Number requires min, max, positive step and mode"
        )
    if minimum > maximum:
        raise WorkbookValidationError(
            f"row {row.row_number}: native_min_value must not exceed native_max_value"
        )
    if step <= 0:
        raise WorkbookValidationError(
            f"row {row.row_number}: native_step must be positive"
        )
    if mode not in NUMBER_MODES:
        raise WorkbookValidationError(
            f"row {row.row_number}: unsupported Number mode {mode!r}"
        )

    # Validate a numeric-literal default exactly against the same range/step
    # contract used by runtime Number writes. Dynamic expressions are checked
    # later by the production compiler and at runtime against this metadata.
    try:
        literal = ast.literal_eval(row["default_source_expression"])
    except (SyntaxError, ValueError):
        return
    if isinstance(literal, bool) or not isinstance(literal, (int, float)):
        return
    try:
        value = Decimal(str(literal))
        low = Decimal(str(minimum))
        high = Decimal(str(maximum))
        increment = Decimal(str(step))
    except InvalidOperation as err:
        raise WorkbookValidationError(
            f"row {row.row_number}: invalid Number decimal metadata"
        ) from err
    if not low <= value <= high:
        raise WorkbookValidationError(
            f"row {row.row_number}: numeric default is outside min/max"
        )
    if (value - low) % increment != 0:
        raise WorkbookValidationError(
            f"row {row.row_number}: numeric default does not match native_step"
        )


def _validate_platform_metadata(row: EntityRow) -> None:
    entity_type = row["entity_type"]
    if entity_type not in ENTITY_TYPES:
        raise WorkbookValidationError(
            f"row {row.row_number}: unsupported entity_type {entity_type!r}"
        )
    key = row["key"]
    translation_key = row["translation_key"]
    if not IDENTIFIER.fullmatch(key):
        raise WorkbookValidationError(f"row {row.row_number}: invalid key {key!r}")
    if not IDENTIFIER.fullmatch(translation_key):
        raise WorkbookValidationError(
            f"row {row.row_number}: invalid translation_key {translation_key!r}"
        )
    if row["icon"] is not None and not ICON.fullmatch(row["icon"]):
        raise WorkbookValidationError(
            f"row {row.row_number}: icon must be an mdi: icon"
        )
    if row["state_class"] not in STATE_CLASSES | {None}:
        raise WorkbookValidationError(
            f"row {row.row_number}: unsupported state_class {row['state_class']!r}"
        )
    if row["entity_category"] not in ENTITY_CATEGORIES | {None}:
        raise WorkbookValidationError(
            f"row {row.row_number}: unsupported entity_category "
            f"{row['entity_category']!r}"
        )

    device_class = row["device_class"]
    allowed_device_classes = {
        "sensor": SENSOR_DEVICE_CLASSES,
        "binary_sensor": BINARY_SENSOR_DEVICE_CLASSES,
        "number": NUMBER_DEVICE_CLASSES,
        "switch": frozenset(),
    }[entity_type]
    if device_class is not None and device_class not in allowed_device_classes:
        raise WorkbookValidationError(
            f"row {row.row_number}: unsupported {entity_type} device_class "
            f"{device_class!r}"
        )

    number_only = (
        "native_min_value", "native_max_value", "native_step", "mode"
    )
    if entity_type == "number":
        if row["state_class"] is not None:
            raise WorkbookValidationError(
                f"row {row.row_number}: Number cannot define state_class"
            )
        _validate_number_metadata(row)
    elif any(row[field] is not None for field in number_only):
        raise WorkbookValidationError(
            f"row {row.row_number}: {entity_type} cannot define Number metadata"
        )

    if entity_type == "binary_sensor":
        forbidden = (
            "native_unit_of_measurement", "state_class",
            "suggested_display_precision",
        )
        if any(row[field] is not None for field in forbidden):
            raise WorkbookValidationError(
                f"row {row.row_number}: binary_sensor has incompatible metadata"
            )

    options = row["options"]
    if row["device_class"] == "enum":
        if entity_type != "sensor" or options is None:
            raise WorkbookValidationError(
                f"row {row.row_number}: enum Sensor requires options"
            )
    elif options is not None:
        raise WorkbookValidationError(
            f"row {row.row_number}: options require the enum Sensor device class"
        )

    if entity_type == "switch":
        forbidden = (
            "device_class", "native_unit_of_measurement", "state_class",
            "suggested_display_precision",
        )
        if any(row[field] is not None for field in forbidden):
            raise WorkbookValidationError(
                f"row {row.row_number}: switch has incompatible metadata"
            )
        try:
            default = ast.literal_eval(row["default_source_expression"])
        except (SyntaxError, ValueError) as err:
            raise WorkbookValidationError(
                f"row {row.row_number}: switch default must be True or False"
            ) from err
        if type(default) is not bool:
            raise WorkbookValidationError(
                f"row {row.row_number}: switch default must be True or False"
            )


def _validate_rounding_precision(row: EntityRow) -> None:
    """Keep every generated round call aligned with display precision."""
    try:
        tree = ast.parse(row["default_source_expression"], mode="eval")
    except SyntaxError:
        # The production compiler reports the more useful syntax error later.
        return

    for node in ast.walk(tree):
        if (
            not isinstance(node, ast.Call)
            or not isinstance(node.func, ast.Name)
            or node.func.id != "round"
        ):
            continue
        if node.keywords or len(node.args) not in {1, 2}:
            raise WorkbookValidationError(
                f"row {row.row_number}: round requires a value and optional "
                "literal precision"
            )
        if len(node.args) == 1:
            digits = 0
        else:
            try:
                digits = ast.literal_eval(node.args[1])
            except (SyntaxError, ValueError) as err:
                raise WorkbookValidationError(
                    f"row {row.row_number}: round precision must be a literal integer"
                ) from err
            if isinstance(digits, bool) or not isinstance(digits, int):
                raise WorkbookValidationError(
                    f"row {row.row_number}: round precision must be a literal integer"
                )
        if not -MAX_ROUND_DIGITS <= digits <= MAX_ROUND_DIGITS:
            raise WorkbookValidationError(
                f"row {row.row_number}: round precision exceeds "
                f"{MAX_ROUND_DIGITS} digits"
            )

        display_precision = row["suggested_display_precision"]
        if display_precision is None or digits != display_precision:
            raise WorkbookValidationError(
                f"row {row.row_number}: round precision {digits} does not match "
                f"suggested_display_precision {display_precision!r}"
            )


def load_translation_files(
    paths: Iterable[Path],
) -> list[tuple[Path, Mapping[str, Any]]]:
    """Load every required translation document."""
    result: list[tuple[Path, Mapping[str, Any]]] = []
    for path in paths:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as err:
            raise WorkbookValidationError(
                f"cannot read translation file {path}: {err}"
            ) from err
        if not isinstance(payload, Mapping):
            raise WorkbookValidationError(
                f"translation file {path} must contain an object"
            )
        result.append((path, payload))
    if not result:
        raise WorkbookValidationError("at least one translation file is required")
    return result


def _validate_translations(
    rows: Sequence[EntityRow],
    translations: Sequence[tuple[Path, Mapping[str, Any]]],
) -> None:
    for path, document in translations:
        root = document.get("entity")
        for row in rows:
            platform = (
                root.get(row["entity_type"])
                if isinstance(root, Mapping)
                else None
            )
            translated = (
                platform.get(row["translation_key"])
                if isinstance(platform, Mapping)
                else None
            )
            name = translated.get("name") if isinstance(translated, Mapping) else None
            if not isinstance(name, str) or not name.strip():
                raise WorkbookValidationError(
                    f"row {row.row_number}: translation {row['entity_type']}."
                    f"{row['translation_key']} is missing from {path}"
                )


def validate_rows(
    rows: Sequence[EntityRow],
    *,
    compile_expressions: Callable[[Mapping[str, str]], Any],
    translations: Sequence[tuple[Path, Mapping[str, Any]]],
) -> None:
    """Validate all row, cross-row, translation and compiler invariants."""
    for attribute in ("key", "unique_id", "entity_id"):
        _ensure_unique(rows, attribute)
    seen_translations: dict[tuple[str, str], int] = {}
    for row in rows:
        _validate_platform_metadata(row)
        _validate_rounding_precision(row)
        identity = (row["entity_type"], row["translation_key"])
        if identity in seen_translations:
            raise WorkbookValidationError(
                f"row {row.row_number}: duplicate platform translation key "
                f"{identity!r}; first used on row {seen_translations[identity]}"
            )
        seen_translations[identity] = row.row_number
    _validate_translations(rows, translations)

    expressions = {
        row["key"]: row["default_source_expression"]
        for row in rows
        if row["default_source_expression"].strip() != "drop()"
    }
    try:
        compile_expressions(expressions)
    except Exception as err:
        raise WorkbookValidationError(
            f"default_source_expression compilation failed: {err}"
        ) from err


def _enum_reference(row: EntityRow, field: str) -> str:
    value = row[field]
    if value is None:
        return "None"
    if field == "device_class":
        prefix = {
            "sensor": "SensorDeviceClass",
            "binary_sensor": "BinarySensorDeviceClass",
            "number": "NumberDeviceClass",
        }[row["entity_type"]]
    elif field == "state_class":
        prefix = "SensorStateClass"
    elif field == "entity_category":
        prefix = "EntityCategory"
    elif field == "mode":
        prefix = "NumberMode"
    elif field == "entity_type":
        prefix = "EntityType"
    else:  # pragma: no cover - guarded by the renderer
        raise AssertionError(field)
    return f"{prefix}.{str(value).upper()}"


def render_descriptions(
    rows: Sequence[EntityRow],
    workbook_name: str,
    *,
    source_sha256: str,
    generator_sha256: str,
    catalog_revision: int,
) -> str:
    """Render deterministic, importable Python after successful validation."""
    lines = [
        '"""Generated static Solesta EntityDescriptions.',
        "",
        "DO NOT EDIT MANUALLY.",
        f"Source: tools/sources/{workbook_name}",
        f"Source SHA-256: {source_sha256}",
        "Generator: tools/generate_thetruth.py",
        f"Generator SHA-256: {generator_sha256}",
        f"Catalog revision: {catalog_revision}",
        '"""',
        "",
        "from homeassistant.components.binary_sensor import BinarySensorDeviceClass",
        "from homeassistant.components.number import NumberDeviceClass, NumberMode",
        "from homeassistant.components.sensor import "
        "SensorDeviceClass, SensorStateClass",
        "from homeassistant.const import EntityCategory",
        "",
        "from .description import Description, EntityType",
        "",
        "",
        "DESCRIPTIONS: list[Description] = [",
    ]
    enum_fields = {
        "device_class", "state_class", "entity_category", "mode", "entity_type"
    }
    for row in rows:
        lines.append("    Description(")
        for field in DESCRIPTION_COLUMNS:
            rendered = (
                _enum_reference(row, field)
                if field in enum_fields
                else repr(row[field])
            )
            lines.append(f"        {field}={rendered},")
        lines.append("    ),")
    lines.extend(("]", ""))
    result = "\n".join(lines)
    ast.parse(result)
    return result


def _validate_catalog(revision: object, rows: Sequence[EntityRow]) -> None:
    if isinstance(revision, bool) or revision != CATALOG_REVISION:
        raise WorkbookValidationError(
            f"catalog_revision must be {CATALOG_REVISION}, got {revision!r}"
        )
    actual_keys = tuple(row["key"] for row in rows)
    if actual_keys != EXPECTED_KEYS:
        missing = sorted(set(EXPECTED_KEYS) - set(actual_keys))
        unexpected = sorted(set(actual_keys) - set(EXPECTED_KEYS))
        details = []
        if missing:
            details.append("missing: " + ", ".join(missing))
        if unexpected:
            details.append("unexpected: " + ", ".join(unexpected))
        if not details:
            details.append("entity keys are not in the canonical order")
        raise WorkbookValidationError(
            "EntityDescriptions does not match catalog revision "
            f"{CATALOG_REVISION} ({'; '.join(details)})"
        )


def load_workbook_rows(path: Path) -> list[EntityRow]:
    """Read the entity worksheet by header name, including formula caches."""
    try:
        from openpyxl import load_workbook
    except ImportError as err:  # pragma: no cover - environment-specific
        raise RuntimeError("generate_thetruth.py requires openpyxl") from err

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise WorkbookValidationError(f"missing worksheet {SHEET_NAME!r}")
        sheet = workbook[SHEET_NAME]
        header_row: int | None = None
        headers: list[str | None] = []
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            candidate = [
                str(value).strip() if value is not None and str(value).strip() else None
                for value in cells
            ]
            if "key" in candidate:
                header_row = row_number
                headers = candidate
                break
        if header_row is None:
            raise WorkbookValidationError(
                "cannot find the EntityDescriptions header row"
            )

        revision: object = None
        for cells in sheet.iter_rows(
            min_row=1, max_row=header_row - 1, values_only=True
        ):
            for index, value in enumerate(cells[:-1]):
                if value == "catalog_revision":
                    if revision is not None:
                        raise WorkbookValidationError(
                            "catalog_revision occurs more than once"
                        )
                    revision = cells[index + 1]

        non_empty_headers = [header for header in headers if header is not None]
        duplicates = sorted(
            {
                header
                for header in non_empty_headers
                if non_empty_headers.count(header) > 1
            }
        )
        if duplicates:
            raise WorkbookValidationError(
                "duplicate worksheet columns: " + ", ".join(duplicates)
            )
        missing = sorted(REQUIRED_COLUMNS - set(non_empty_headers))
        if missing:
            raise WorkbookValidationError(
                "missing required columns: " + ", ".join(missing)
            )
        unsupported = sorted(
            set(non_empty_headers) - REQUIRED_COLUMNS - AUDIT_COLUMNS
        )
        if unsupported:
            raise WorkbookValidationError(
            "unsupported EntityDescriptions columns: " + ", ".join(unsupported)
            )

        raw_rows: list[tuple[int, Mapping[str, Any]]] = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            header_row + 1,
        ):
            raw_rows.append(
                (
                    row_number,
                    {
                        header: cells[index] if index < len(cells) else None
                        for index, header in enumerate(headers)
                        if header is not None
                    },
                )
            )
        rows = normalize_rows(raw_rows)
        _validate_catalog(revision, rows)
        return rows
    finally:
        workbook.close()


def load_production_compiler(
    integration_dir: Path,
) -> Callable[[Mapping[str, str]], Any]:
    """Load ``compile.py`` and its pure helper without importing Home Assistant."""
    package_name = "_solesta_source_generator_runtime"
    package = ModuleType(package_name)
    package.__path__ = [str(integration_dir)]
    sys.modules[package_name] = package
    for name in ("expression_functions", "compile"):
        full_name = f"{package_name}.{name}"
        spec = importlib.util.spec_from_file_location(
            full_name, integration_dir / f"{name}.py"
        )
        if spec is None or spec.loader is None:
            raise RuntimeError(f"cannot load production module {name}")
        module = importlib.util.module_from_spec(spec)
        sys.modules[full_name] = module
        spec.loader.exec_module(module)
    return sys.modules[f"{package_name}.compile"].compile_source_expressions


def generate_text(
    workbook_path: Path,
    translation_paths: Sequence[Path],
    *,
    integration_dir: Path,
    generator_path: Path | None = None,
) -> str:
    """Validate the canonical input and return deterministic Python source."""
    rows = load_workbook_rows(workbook_path)
    translations = load_translation_files(translation_paths)
    compiler = load_production_compiler(integration_dir)
    validate_rows(
        rows,
        compile_expressions=compiler,
        translations=translations,
    )
    source_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    actual_generator_path = generator_path or Path(__file__)
    generator_hash = hashlib.sha256(actual_generator_path.read_bytes()).hexdigest()
    return render_descriptions(
        rows,
        workbook_path.name,
        source_sha256=source_hash,
        generator_sha256=generator_hash,
        catalog_revision=CATALOG_REVISION,
    )


def generate(
    workbook_path: Path,
    output_path: Path,
    translation_paths: Sequence[Path],
    *,
    integration_dir: Path,
    generator_path: Path | None = None,
) -> None:
    """Validate first and atomically replace the generated output."""
    rendered = generate_text(
        workbook_path,
        translation_paths,
        integration_dir=integration_dir,
        generator_path=generator_path,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8",
            dir=output_path.parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            handle.write(rendered)
            temporary_path = Path(handle.name)
        temporary_path.replace(output_path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def _default_translation_paths(integration_dir: Path) -> list[Path]:
    paths = [integration_dir / "strings.json"]
    paths.extend(sorted((integration_dir / "translations").glob("*.json")))
    return paths


def main(argv: Sequence[str] | None = None) -> int:
    tools_dir = Path(__file__).resolve().parent
    integration_dir = tools_dir.parent / "custom_components" / "solesta"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=tools_dir / "sources" / "EntityDescriptions.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=integration_dir / "thetruth.py",
    )
    parser.add_argument(
        "--translation",
        action="append",
        type=Path,
        dest="translations",
        help="translation JSON to validate; repeat for multiple locales",
    )
    parser.add_argument(
        "--integration-dir",
        type=Path,
        default=integration_dir,
        help=(
            "integration directory containing compile.py and translations; "
            "defaults to custom_components/solesta"
        ),
    )
    args = parser.parse_args(argv)
    translation_paths = (
        args.translations
        if args.translations is not None
        else _default_translation_paths(args.integration_dir)
    )
    try:
        generate(
            args.workbook,
            args.output,
            translation_paths,
            integration_dir=args.integration_dir,
        )
    except (OSError, RuntimeError, WorkbookValidationError) as err:
        parser.error(str(err))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
